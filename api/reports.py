"""
Отчёты: отчёт о занятии (PDF), сертификат (PDF), выгрузка (Excel).

ТЗ: отчёт о практическом занятии с действиями, замечаниями, временем
заполнения карточки, отклонением от норматива и грамматикой.
Экспорт в PDF и Excel. Сертификаты в PDF.
"""
from __future__ import annotations

import io
import os
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from . import config as C

_FONT = "Helvetica"
for d in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu"):
    if os.path.exists(f"{d}/DejaVuSans.ttf"):
        pdfmetrics.registerFont(TTFont("DV", f"{d}/DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DVB", f"{d}/DejaVuSans-Bold.ttf"))
        _FONT = "DV"
        break
_BOLD = "DVB" if _FONT == "DV" else "Helvetica-Bold"

ST = {
    "h1": ParagraphStyle("h1", fontName=_BOLD, fontSize=16, leading=20, spaceAfter=6),
    "h2": ParagraphStyle("h2", fontName=_BOLD, fontSize=11, leading=14,
                         spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#0369a1")),
    "p": ParagraphStyle("p", fontName=_FONT, fontSize=9.5, leading=13),
    "s": ParagraphStyle("s", fontName=_FONT, fontSize=8, leading=10,
                        textColor=colors.HexColor("#64748b")),
}


def _tbl(rows, widths, head=True):
    t = Table([[Paragraph(str(c), ST["p"]) for c in r] for r in rows], colWidths=widths)
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]
    if head:
        style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e0f2fe")))
    t.setStyle(TableStyle(style))
    return t


def otchet_zanyatiya(sess: dict, user: dict, sc: dict) -> bytes:
    it = sess["itog"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Отчёт о практическом занятии")
    E = [Paragraph("Отчёт о практическом занятии", ST["h1"]),
         Paragraph("Тренажёр оператора ДДС · Департамент ГОЧСиПБ города Москвы", ST["s"]),
         Spacer(1, 6)]

    ball = sess.get("override_ball") if sess.get("override_ball") is not None else it["ball"]
    E.append(_tbl([
        ["Обучающийся", user["name"], "Дата", time.strftime("%d.%m.%Y %H:%M",
                                                             time.localtime(sess["started"]))],
        ["Сценарий", sc.get("nazvanie", sess["scenario_id"]), "Балл",
         f"{ball} из 100 — {it['verdikt']}"],
        ["Время карточки", f"{it['t_sec']} с", "Норматив",
         f"{C.NORM['kartochka_sec']} с · отклонение {round(it['t_sec'] - C.NORM['kartochka_sec'], 1):+} с"],
    ], [34 * mm, 58 * mm, 26 * mm, 60 * mm], head=False))

    if sess.get("override_ball") is not None:
        E.append(Spacer(1, 4))
        E.append(Paragraph(f"Оценка изменена преподавателем: {it['ball']} → "
                           f"{sess['override_ball']}. Причина: {sess.get('override_reason') or '—'}",
                           ST["s"]))

    E.append(Paragraph("Граф доказательств", ST["h2"]))
    rows = [["Факт", "Результат", "Вес", "Основание"]]
    for f in it["fakty"]:
        rows.append([f["nazvanie"], "зачтено" if f["proyden"] else "НЕ ЗАЧТЕНО",
                     f["ves"], f["istochnik"]])
    E.append(_tbl(rows, [62 * mm, 26 * mm, 12 * mm, 78 * mm]))

    zam = []
    for f in it["fakty"]:
        d = f["detali"]
        if f["kod"] == "adres" and not f["proyden"]:
            zam.append(f"Адрес не доведён до эталона. Введено: «{d.get('vvod', '')}». "
                       f"Не хватает: {', '.join(d.get('propushcheno', []))}.")
        if f["kod"] == "sluzhby" and d.get("propushcheno"):
            zam.append(f"Не назначены службы: {', '.join(d['propushcheno'])}.")
        if f["kod"] == "polnota" and d.get("ne_zapolneno"):
            zam.append(f"Не заполнены поля: {', '.join(d['ne_zapolneno'])}.")
        if f["kod"] == "rech":
            for n in d.get("spisok", []):
                zam.append(f"Речь: «{n['fragment']}» — {n['opisanie']}"
                           + (f". Рекомендуется: «{n['zamena']}»." if n.get("zamena") else "."))
        if f["kod"] == "grammatika":
            for g in d.get("spisok", []):
                zam.append(f"Грамматика: «{g['fragment']}» — {g['zamechanie']}.")
    E.append(Paragraph("Замечания", ST["h2"]))
    for z in zam or ["Замечаний нет."]:
        E.append(Paragraph("• " + z, ST["p"]))

    zv = sess["data"].get("zvonki") or []
    if zv:
        E.append(Paragraph("Переговоры по IP-телефону", ST["h2"]))
        for c in zv:
            E.append(Paragraph(f"<b>{'Входящий' if c.get('vhodyashchiy') else 'Исходящий'}: "
                               f"{c['kontakt']['nazvanie']}</b> — "
                               f"{'ответил' if c['ishod'] == 'otvetil' else 'не ответил'}", ST["p"]))
            for m in c.get("dialog", []):
                kto = "Диспетчер" if m["kto"] == "dispetcher" else "Абонент"
                E.append(Paragraph(f"&nbsp;&nbsp;{kto}: {m['tekst']}", ST["p"]))
        st = sess["data"].get("statusy") or []
        if st:
            E.append(Paragraph("Статусы службы", ST["h2"]))
            t0 = sess["data"]["nachalo"]
            for x in st:
                E.append(Paragraph(f"{round(x['t'] - t0, 1)} с · <b>{x['nazvanie']}</b>"
                                   + (f" — {x['kommentariy']}" if x.get("kommentariy") else ""), ST["p"]))
    E.append(Paragraph("Ход разговора", ST["h2"]))
    for m in sess["data"].get("dialog", []):
        kto = "Оператор" if m["kto"] == "operator" else "Заявитель"
        E.append(Paragraph(f"<b>{round(m['t_ms'] / 1000, 1)} с · {kto}:</b> {m['tekst']}",
                           ST["p"]))
    doc.build(E)
    return buf.getvalue()


def sertifikat(user: dict, stat: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=25 * mm,
                            rightMargin=25 * mm, topMargin=25 * mm, bottomMargin=20 * mm,
                            title="Сертификат")
    big = ParagraphStyle("b", fontName=_BOLD, fontSize=30, leading=36, alignment=1)
    mid = ParagraphStyle("m", fontName=_FONT, fontSize=14, leading=20, alignment=1)
    nm = ParagraphStyle("n", fontName=_BOLD, fontSize=22, leading=28, alignment=1,
                        textColor=colors.HexColor("#0369a1"))
    E = [Paragraph("СЕРТИФИКАТ", big), Spacer(1, 10),
         Paragraph("о прохождении практической подготовки на тренажёре оператора ДДС", mid),
         Spacer(1, 22), Paragraph(user["name"], nm), Spacer(1, 18),
         Paragraph(f"Выполнено вызовов: {stat['n']} · средний балл: {stat['ball']} · "
                   f"в норматив {C.NORM['kartochka_sec']} с: {stat['v_norm']}%", mid),
         Spacer(1, 8),
         Paragraph(f"Освоение навыков по модели BKT: {stat['osvoeno']} из {stat['navykov']}",
                   mid),
         Spacer(1, 40),
         Paragraph(f"Департамент ГОЧСиПБ города Москвы · {time.strftime('%d.%m.%Y')}",
                   ST["s"])]
    doc.build(E)
    return buf.getvalue()


def excel(sessii: list[dict], users: dict[int, dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Занятия"
    head = ["Дата", "Обучающийся", "Сценарий", "Балл", "Итоговый балл",
            "Время, с", "В норматив", "Адрес", "Тип", "Службы", "Речь", "Грамматика"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E0F2FE")
    for s in sessii:
        it = s["itog"]
        fk = {f["kod"]: f["proyden"] for f in it["fakty"]}
        ws.append([
            time.strftime("%d.%m.%Y %H:%M", time.localtime(s["started"])),
            users.get(s["user_id"], {}).get("name", "?"), s["scenario_id"],
            it["ball"], s.get("override_ball") if s.get("override_ball") is not None else it["ball"],
            it["t_sec"], "да" if it["v_normativ"] else "нет",
            *["да" if fk.get(k) else "нет" for k in
              ("adres", "tip", "sluzhby", "rech", "grammatika")],
        ])
    for col, w in zip("ABCDEFGHIJKL", (17, 26, 14, 7, 12, 9, 11, 8, 8, 9, 8, 11)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")

    ws2 = wb.create_sheet("Форма 1-112")
    ws2.append(["№", "Наименование", "Требование", "Факт"])
    ts = [s["itog"]["t_sec"] for s in sessii] or [0]
    ws2.append([3, "Среднее время опроса до доступности карточки, сек",
                C.NORM["kartochka_sec"], round(sum(ts) / len(ts), 1)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
