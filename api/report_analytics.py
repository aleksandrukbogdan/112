"""Exports consume the exact frozen payload displayed in the browser."""
import io
import json
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.graphics.shapes import Drawing, Line, Circle, String
from .reports import ST, _tbl


def excel(a):
    wb=Workbook();ws=wb.active;ws.title='Параметры'
    for row in [('Снимок',a['snapshot_id']),('Рассчитан UTC',a['created']),('Фильтры',json.dumps(a['filters'],ensure_ascii=False)),('Попыток',a['n']),('Единица',a['unit'])]:ws.append(row)
    for k,v in a['summary'].items():ws.append([k,v])
    for name,columns,rows in [
        ('Попытки',['ID','Ученик','Режим','Сценарий','Сложность','Балл','Зачёт','Критические ошибки','Секунд','Помощь','Рубрика','Повтор'],
         [[r['id'],r['name'],r['kind'],r['scenario'],r['difficulty'],r['ball'],r['passed'],', '.join(r['critical_errors']),r['t_sec'],r['assisted'],r['rubric'],r['replay']] for r in a['attempts']]),
        ('Навыки',['Программа','Навык','Применимых проверок','Успешных','Доля','Не оценено'],
         [[r['kind'],r['name'],r['n'],r['success'],r['rate'],r['unknown']] for r in a['skills']]),
        ('Ошибки',['Критерий','Ошибок','Применимых','Доля','Критический'],
         [[r['name'],r['errors'],r['n'],r['rate'],r['critical']] for r in a['errors']]),
        ('Время',['Метрика','ID попытки','Секунд','Лимит'],
         [[k,p['id'],p['value'],p['limit']] for k,r in a['timings'].items() for p in r['points']]),
        ('Помощник',['Показатель','Значение'],[[k,str(v)] for k,v in a['helper'].items()]),
        ('Прогноз',['Версия','n','Brier','Brier частоты','Статус'],
         [[x['model'],x['n'],x['brier'],x['baseline_brier'],x['status']] for x in a['forecast']['by_model']]),
    ]:
        ws=wb.create_sheet(name);ws.append(columns)
        for row in rows:
            # Prevent formula execution when source text is opened in Excel.
            ws.append([("'"+v if isinstance(v,str) and v[:1] in ('=','+','-','@') else v) for v in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='153D60')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in col)+2))
    buf=io.BytesIO();wb.save(buf);return buf.getvalue()


def pdf(a):
    buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=16*mm,rightMargin=16*mm,topMargin=13*mm,bottomMargin=13*mm)
    elems=[Paragraph('Учебная аналитика 112 / ДДС',ST['h1']),Paragraph('Снимок: '+escape(a['snapshot_id']),ST['s']),
           Paragraph('Фильтры: '+escape(json.dumps(a['filters'],ensure_ascii=False)),ST['s']),Spacer(1,10)]
    s=a['summary'];elems.append(Paragraph(f"Завершено: {s['completed']} · зачтено: {s['passed']} · попыток с критическими ошибками: {s['critical_attempts']} · остановлено преподавателем: {s['stopped']}",ST['p']))
    if a['attempts']:
        d=Drawing(690,130);d.add(Line(30,15,675,15));d.add(Line(30,15,30,120))
        thresholds={r.get('pass_threshold',70) for r in a['attempts']}
        if len(thresholds)==1:d.add(Line(30,15+next(iter(thresholds)),675,15+next(iter(thresholds)),strokeColor=__import__('reportlab.lib.colors',fromlist=['HexColor']).HexColor('#aaaaaa')))
        n=len(a['attempts'])
        for i,r in enumerate(a['attempts']):
            x=30+i*635/max(n-1,1);y=15+r['ball']
            from reportlab.lib.colors import HexColor
            d.add(Circle(x,y,3,fillColor=HexColor('#19734b' if r['passed'] else '#bd3636'),strokeColor=None))
        elems.extend([Paragraph('Баллы по попыткам. Цвет обозначает зачёт. Линия порога показана только при общем пороге всех попыток.',ST['s']),d])
    elems.append(Paragraph('Навыки: доля успешных применимых проверок',ST['h2']))
    rows=[['Программа','Навык','Успех / применимо','Доля','Не оценено']]+[[r['kind'],r['name'],f"{r['success']} / {r['n']}",f"{100*r['rate']:.1f}%" if r['rate'] is not None else '—',r['unknown']] for r in a['skills']]
    elems.append(_tbl(rows,[25*mm,95*mm,55*mm,35*mm,40*mm]))
    elems.append(Paragraph('Ошибки',ST['h2']))
    rows=[['Критерий','Ошибки / применимо','Критический']]+[[r['name'],f"{r['errors']} / {r['n']}",'Да' if r['critical'] else 'Нет'] for r in a['errors']]
    elems.append(_tbl(rows,[155*mm,60*mm,35*mm]))
    elems.append(Paragraph('Качество помощника и прогноза',ST['h2']))
    h=a['helper'];elems.append(Paragraph(f"Предложений показано: {h['shown']}; проверено преподавателем: {h['reviewed']}; правильных: {h['correct']}; ошибочных: {h['wrong']}. Принятие учеником не доказывает правильность.",ST['p']))
    for m in a['forecast']['by_model']:elems.append(Paragraph(escape(f"{m['model']}: n={m['n']}, Brier={m['brier']:.4f}; статус: {m['status']}. Описательная проверка."),ST['p']))
    for w in a['warnings']:elems.append(Paragraph(escape(w),ST['s']))
    elems.append(Paragraph('Попытки',ST['h2']))
    rows=[['Ученик','Программа / сценарий','Балл','Зачёт','Помощь','Время, с']]+[[r['name'],r['kind']+' / '+r['scenario'],r['ball'],'Да' if r['passed'] else 'Нет','Да' if r['assisted'] else 'Нет',r['t_sec']] for r in a['attempts']]
    elems.append(_tbl(rows,[55*mm,90*mm,20*mm,25*mm,25*mm,35*mm]))
    doc.build(elems);return buf.getvalue()
