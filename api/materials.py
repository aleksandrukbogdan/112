"""Versioned teaching sources. Local text extraction, no external document services."""
import hashlib
import io
import json
import time
import uuid
import zipfile
from pathlib import Path
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from . import db, access, session_store as store
from .auth import current_user, role

router=APIRouter(prefix='/api/materials',tags=['materials'])


def extract(filename,body):
    ext=Path(filename).suffix.lower()
    if ext in ('.txt','.md'):
        try:return body.decode('utf-8-sig')
        except UnicodeDecodeError:return body.decode('cp1251')
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(body))
        if len(reader.pages)>300:raise ValueError('Не более 300 страниц')
        return '\n\n'.join(p.extract_text() or '' for p in reader.pages)
    if ext=='.docx':
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            if sum(x.file_size for x in archive.infolist())>100*1024*1024:raise ValueError('Слишком большой распакованный документ')
        from docx import Document
        d=Document(io.BytesIO(body))
        return '\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])
    raise ValueError('Поддерживаются PDF с текстом, DOCX, TXT и Markdown')


@router.post('')
async def upload(file:UploadFile=File(...),title:str=Form(''),u=Depends(role('teacher'))):
    body=await file.read(20*1024*1024+1)
    if len(body)>20*1024*1024:raise HTTPException(413,'Максимум 20 МБ')
    filename=Path(file.filename or 'material.txt').name
    try:text=extract(filename,body).strip()
    except Exception as exc:raise HTTPException(422,'Не удалось извлечь текст: '+str(exc)[:180])
    if not text:raise HTTPException(422,'Текста нет. Для скана сначала выполните OCR и загрузите текст или DOCX.')
    if len(text)>1000000:raise HTTPException(413,'Извлечённый текст больше 1 млн символов')
    mid=uuid.uuid4().hex;digest=hashlib.sha256(body).hexdigest()
    meta={'bytes':len(body),'chars':len(text),'extraction':'local-text','source_version':digest,'has_original_binary':False}
    db.ex('INSERT INTO materials(id,owner_id,title,filename,hash,content,meta,created) VALUES(?,?,?,?,?,?,?,?)',
          (mid,u['id'],title.strip() or filename,filename,digest,text,store.dumps(meta),time.time()))
    db.audit(u,'material_uploaded',{'id':mid,'hash':digest,'filename':filename})
    return {'id':mid,'title':title.strip() or filename,'meta':meta}


def permitted(u):
    if u['role']=='admin':return None
    if u['role']=='teacher':return {x['id'] for x in db.q('SELECT id FROM materials WHERE owner_id=?',(u['id'],))}
    out=set()
    for row in db.q('SELECT config FROM lessons WHERE group_id=?',(u.get('group_id'),)):
        out.update(json.loads(row['config']).get('material_ids',[]))
    return out


@router.get('')
def listing(u=Depends(current_user)):
    ids=permitted(u)
    return [{**r,'meta':json.loads(r['meta'])} for r in db.q('SELECT id,title,filename,hash,meta,created,owner_id FROM materials ORDER BY created DESC') if ids is None or r['id'] in ids]


@router.get('/{mid}')
def read(mid:str,u=Depends(current_user)):
    ids=permitted(u)
    if ids is not None and mid not in ids:raise HTTPException(403,'Материал не назначен')
    row=db.q1('SELECT * FROM materials WHERE id=?',(mid,))
    if not row:raise HTTPException(404,'Материал не найден')
    row['meta']=json.loads(row['meta']);return row
